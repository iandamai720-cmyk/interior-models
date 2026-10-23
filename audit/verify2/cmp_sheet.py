import openpyxl,csv,collections,sys
f=sys.argv[1]
R=list(csv.reader(open('live/assets.csv',encoding='utf-8-sig')));H=R[0]
wb=openpyxl.load_workbook(f,read_only=True);print(wb.sheetnames)
X=[[('' if c is None else str(c)) for c in r] for r in wb['Wszystkie'].iter_rows(values_only=True)]
X=[r for r in X if any(r)]
print('rows',len(X)-1,'hdr',X[0][:len(H)]==H)
d=collections.Counter();ex={}
for a,b in zip(R[1:],X[1:]):
    b=b+['']*(len(a)-len(b))
    for k,(u,v) in enumerate(zip(a,b)):
        if u!=v:
            try:
                if abs(float(u)-float(v))<1e-9: continue
            except: pass
            if u.strip()==v.strip(): continue
            d[H[k]]+=1;ex.setdefault(H[k],(a[H.index('id')],u[:80],v[:80]))
print(dict(d));print(ex)
