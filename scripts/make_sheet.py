"""sklepy.xlsx + assets.csv: every model / store product with store, price, links, model and preview URLs."""
import json, csv, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
ROOT = "/workspace/interior-lib"
A = json.load(open(f"{ROOT}/catalog.json"))["assets"]; BY = {a["id"]: a for a in A}
COLS = ["name", "opis", "category", "room", "store", "product_url", "price_pln", "product_image_url", "preview_url", "model_file_url", "has_own_model",
        "similar_label", "similar_model_id", "similar_model_name", "similar_model_url", "similar_preview_url", "colour", "material", "dimensions_cm",
        "id", "type", "style", "url_uk", "source", "source_title", "license", "shard"]
def dimtxt(a):
    d = a.get("dimensions_cm")
    if d: return " x ".join(f"{d[k]:g}" for k in ("w", "d", "h") if d.get(k))
    d = a.get("dimensions_m")
    return " x ".join(f"{d[k]*100:.0f}" for k in ("x", "y", "z")) if d else ""
rows = []
for a in A:
    if a["type"] == "layout": continue
    rows.append(dict(name=a.get("full_name") or a["name"], category=a["category"], room=a["room"], store=a.get("store") or a.get("source"),
        product_url=a.get("url_pl") or a.get("source_url"), price_pln=a.get("price_pln"), model_file_url=(a.get("model_url") or a.get("url") or "") if a["type"] != "store_product" else "",
        product_image_url=a.get("product_image_url") or "", opis=a.get("opis") or "",
        preview_url=a.get("preview_url") or "", has_own_model=bool(a.get("has_own_model", True)), similar_label=a.get("similar_label") or "", similar_model_id=a.get("similar_model_id") or "",
        similar_model_name=a.get("similar_model_name") or "", similar_model_url=a.get("similar_model_url") or "", similar_preview_url=a.get("similar_preview_url") or "",
        colour=", ".join(a.get("colours") or ([a["colour"]] if a.get("colour") else [])), material=a.get("material") or "", dimensions_cm=dimtxt(a),
        id=a["id"], type=a["type"], style=a["style"], url_uk=a.get("url_uk") or "", source=a.get("source"), source_title=a.get("source_title") or "", license=a.get("license"), shard=a.get("shard") or ""))
with open(f"{ROOT}/assets.csv", "w", newline="", encoding="utf-8-sig") as f:
    w = csv.DictWriter(f, COLS, quoting=csv.QUOTE_MINIMAL); w.writeheader(); w.writerows(rows)
with open(f"{ROOT}/assets_pl.csv", "w", newline="", encoding="utf-8-sig") as f:  # Polish Excel: ';' separator, decimal comma
    w = csv.DictWriter(f, COLS, delimiter=";", quoting=csv.QUOTE_MINIMAL); w.writeheader()
    w.writerows([{k: (str(v).replace(".", ",") if k == "price_pln" and v is not None else v) for k, v in r.items()} for r in rows])
wb = Workbook(); hdr = Font(bold=True, color="FFFFFF"); fill = PatternFill("solid", fgColor="0058A3")
def sheet(ws, rs, cols):
    ws.append(cols)
    for c in ws[1]: c.font, c.fill = hdr, fill
    for r in rs: ws.append([r.get(c) for c in cols])
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for i, c in enumerate(cols, 1): ws.column_dimensions[get_column_letter(i)].width = {"name": 55, "product_url": 50, "model_file_url": 50, "preview_url": 40, "license": 40, "opis": 70, "product_image_url": 40, "similar_model_name": 40}.get(c, 16)
ws = wb.active; ws.title = "Wszystkie"; sheet(ws, rows, COLS)
by = collections.defaultdict(list)
for r in rows:
    if r["store"]: by[r["store"]].append(r)
for s, rs in sorted(by.items()):
    sheet(wb.create_sheet(s[:31].replace("/", "-")), rs, COLS)
ss = wb.create_sheet("Podsumowanie", 0); ss.append(["store/source", "items", "own 3D model", "store product (no own model)", "with price PLN", "avg price PLN"])
for c in ss[1]: c.font, c.fill = hdr, fill
for s, rs in sorted(by.items()):
    pr = [r["price_pln"] for r in rs if r["price_pln"] is not None]
    ss.append([s, len(rs), sum(r["has_own_model"] for r in rs), sum(not r["has_own_model"] for r in rs), len(pr), round(sum(pr) / len(pr)) if pr else None])
ss.column_dimensions["A"].width = 28
wb.save(f"{ROOT}/sklepy.xlsx"); print(len(rows), {s: len(v) for s, v in by.items()})
